import io
import re
from pathlib import Path
from pypdf import PdfReader
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

SKILLS = ["python","java","c++","sql","mysql","postgresql","mongodb","excel","power bi","tableau","pandas","numpy","scikit-learn","machine learning","deep learning","natural language processing","nlp","tensorflow","pytorch","keras","data analysis","data visualization","statistics","matplotlib","seaborn","opencv","computer vision","html","css","javascript","react","flask","django","fastapi","streamlit","git","github","linux","aws","azure","docker","rest api","api","communication","project management","problem solving","data structures","algorithms","jupyter","r programming","spark","hadoop","etl","powerpoint"]
EDUCATION_TERMS = ["bachelor","b.tech","b.e.","bsc","b.sc","master","m.tech","mca","bca","computer science","information technology","degree","diploma","phd"]
EXPERIENCE_PATTERNS = [r"\b\d+\+?\s+years?\s+(?:of\s+)?experience\b",r"\bexperience\s*:\s*[^\n]{1,80}",r"\bintern(?:ship)?\b",r"\bprofessional experience\b",r"\bwork experience\b"]

def extract_text(data: bytes, filename: str) -> str:
    ext=Path(filename).suffix.lower()
    if ext==".pdf":
        reader=PdfReader(io.BytesIO(data))
        text="\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext==".docx":
        doc=Document(io.BytesIO(data))
        parts=[p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows: parts.append(" | ".join(cell.text for cell in row.cells))
        text="\n".join(parts)
    else: raise ValueError("Unsupported file format. Upload PDF or DOCX.")
    return re.sub(r"[ \t]+"," ",text).strip()

def _normalize(text: str) -> str:
    return re.sub(r"\s+"," ",text.lower()).strip()

def detect_skills(text: str) -> list[str]:
    normalized=_normalize(text); found=[]
    for skill in SKILLS:
        pattern=r"(?<![a-z0-9])"+re.escape(skill.lower())+r"(?![a-z0-9])"
        if re.search(pattern,normalized): found.append(skill)
    return sorted(set(found))

def _phrases(text: str, terms: list[str]) -> list[str]:
    low=_normalize(text)
    return [term for term in terms if re.search(r"(?<![a-z0-9])"+re.escape(term)+r"(?![a-z0-9])",low)]

def analyze_resumes(job_description: str, candidates: list[dict]) -> list[dict]:
    if not candidates: return []
    docs=[job_description]+[c["text"] for c in candidates]
    vectorizer=TfidfVectorizer(stop_words="english",ngram_range=(1,2),sublinear_tf=True)
    matrix=vectorizer.fit_transform(docs)
    sims=cosine_similarity(matrix[0:1],matrix[1:]).flatten()
    job_skills=set(detect_skills(job_description)); results=[]
    for candidate,sim in zip(candidates,sims):
        resume_skills=set(detect_skills(candidate["text"]))
        matched=sorted(job_skills & resume_skills); missing=sorted(job_skills-resume_skills)
        coverage=(len(matched)/len(job_skills)*100) if job_skills else 0.0
        similarity=float(sim*100); score=0.65*similarity+0.35*coverage
        text=candidate["text"]; education=_phrases(text,EDUCATION_TERMS); experience=[]
        for pat in EXPERIENCE_PATTERNS: experience.extend(re.findall(pat,text,flags=re.I))
        results.append({"name":candidate["name"],"similarity":similarity,"skill_coverage":coverage,"match_score":score,"matched_skills":matched,"missing_skills":missing,"education":sorted(set(education)),"experience":sorted(set(x.strip() for x in experience))})
    return sorted(results,key=lambda r:r["match_score"],reverse=True)
